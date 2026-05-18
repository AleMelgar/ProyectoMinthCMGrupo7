import io
from datetime import datetime

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


def generate_pdf_report(
    form_name: str,
    period_name: str,
    employees_data: list[dict],
) -> io.BytesIO:
    """Genera un reporte PDF con los resultados de evaluación.

    employees_data: lista de dicts con keys:
        employee_name, goals (list), kpis (list), puntaje_total, comentarios
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=0.5 * inch, bottomMargin=0.5 * inch)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle("CustomTitle", parent=styles["Title"], fontSize=18, spaceAfter=6)
    subtitle_style = ParagraphStyle("Subtitle", parent=styles["Normal"], fontSize=12, textColor=colors.grey, spaceAfter=20)
    heading_style = ParagraphStyle("Heading", parent=styles["Heading2"], fontSize=14, spaceAfter=8)

    story.append(Paragraph("PerformTrack — Reporte de Evaluación", title_style))
    story.append(Paragraph(f"Formulario: {form_name} | Período: {period_name}", subtitle_style))
    story.append(Paragraph(f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]))
    story.append(Spacer(1, 20))

    if not employees_data:
        story.append(Paragraph("No hay datos de evaluación disponibles.", styles["Normal"]))
        doc.build(story)
        buffer.seek(0)
        return buffer

    # Tabla resumen
    story.append(Paragraph("Resumen General", heading_style))
    summary_header = ["Empleado", "Puntaje Total", "Estado"]
    summary_rows = [summary_header]
    for emp in employees_data:
        summary_rows.append([
            emp.get("employee_name", "—"),
            f"{emp.get('puntaje_total', 0):.1f}%",
            emp.get("estado", "—"),
        ])

    summary_table = Table(summary_rows, colWidths=[3 * inch, 1.5 * inch, 1.5 * inch])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1565C0")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 10),
        ("ALIGN", (1, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.white]),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 20))

    # Detalle por empleado
    for emp in employees_data:
        story.append(Paragraph(f"Empleado: {emp.get('employee_name', '—')}", heading_style))

        # OKRs
        goals = emp.get("goals", [])
        if goals:
            story.append(Paragraph("Objetivos OKR", styles["Heading3"]))
            goal_header = ["Objetivo", "Peso", "Progreso"]
            goal_rows = [goal_header]
            for g in goals:
                goal_rows.append([
                    g.get("objetivo_okr", "—"),
                    f"{g.get('peso', 0):.0f}%",
                    f"{g.get('progreso', 0):.1f}%",
                ])
            goal_table = Table(goal_rows, colWidths=[3.5 * inch, 1 * inch, 1 * inch])
            goal_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2196F3")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(goal_table)
            story.append(Spacer(1, 10))

        # KPIs
        kpis = emp.get("kpis", [])
        if kpis:
            story.append(Paragraph("KPIs", styles["Heading3"]))
            kpi_header = ["KPI", "Actual", "Meta", "Porcentaje"]
            kpi_rows = [kpi_header]
            for k in kpis:
                kpi_rows.append([
                    k.get("kpi_nombre", "—"),
                    f"{k.get('valor_actual', 0):.1f}",
                    f"{k.get('valor_meta', 0):.1f}",
                    f"{k.get('porcentaje', 0):.1f}%",
                ])
            kpi_table = Table(kpi_rows, colWidths=[2.5 * inch, 1 * inch, 1 * inch, 1 * inch])
            kpi_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4CAF50")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ALIGN", (1, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.lightgrey),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(kpi_table)
            story.append(Spacer(1, 10))

        if emp.get("comentarios"):
            story.append(Paragraph(f"Comentarios: {emp['comentarios']}", styles["Normal"]))

        story.append(Spacer(1, 20))

    doc.build(story)
    buffer.seek(0)
    return buffer


def generate_excel_report(
    form_name: str,
    period_name: str,
    employees_data: list[dict],
) -> io.BytesIO:
    """Genera un reporte Excel con los resultados de evaluación."""
    wb = openpyxl.Workbook()

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1565C0", end_color="1565C0", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    # --- Hoja: Resumen ---
    ws = wb.active
    ws.title = "Resumen"
    ws.append(["PerformTrack — Reporte de Evaluación"])
    ws.merge_cells("A1:D1")
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Formulario: {form_name}", f"Período: {period_name}"])
    ws.append([f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}"])
    ws.append([])

    headers = ["Empleado", "Puntaje Total", "Estado", "Comentarios"]
    ws.append(headers)
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=5, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for emp in employees_data:
        ws.append([
            emp.get("employee_name", "—"),
            round(emp.get("puntaje_total", 0), 1),
            emp.get("estado", "—"),
            emp.get("comentarios", ""),
        ])

    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 15
    ws.column_dimensions["D"].width = 40

    # --- Hoja: OKRs ---
    ws_goals = wb.create_sheet("OKRs")
    goal_headers = ["Empleado", "Objetivo OKR", "Peso (%)", "Progreso (%)"]
    ws_goals.append(goal_headers)
    for col_idx, _ in enumerate(goal_headers, 1):
        cell = ws_goals.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = PatternFill(start_color="2196F3", end_color="2196F3", fill_type="solid")
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for emp in employees_data:
        for g in emp.get("goals", []):
            ws_goals.append([
                emp.get("employee_name", "—"),
                g.get("objetivo_okr", "—"),
                round(g.get("peso", 0), 1),
                round(g.get("progreso", 0), 1),
            ])

    ws_goals.column_dimensions["A"].width = 25
    ws_goals.column_dimensions["B"].width = 45
    ws_goals.column_dimensions["C"].width = 12
    ws_goals.column_dimensions["D"].width = 14

    # --- Hoja: KPIs ---
    ws_kpis = wb.create_sheet("KPIs")
    kpi_headers = ["Empleado", "KPI", "Valor Actual", "Valor Meta", "Porcentaje (%)"]
    ws_kpis.append(kpi_headers)
    for col_idx, _ in enumerate(kpi_headers, 1):
        cell = ws_kpis.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center")

    for emp in employees_data:
        for k in emp.get("kpis", []):
            ws_kpis.append([
                emp.get("employee_name", "—"),
                k.get("kpi_nombre", "—"),
                round(k.get("valor_actual", 0), 1),
                round(k.get("valor_meta", 0), 1),
                round(k.get("porcentaje", 0), 1),
            ])

    ws_kpis.column_dimensions["A"].width = 25
    ws_kpis.column_dimensions["B"].width = 30
    ws_kpis.column_dimensions["C"].width = 14
    ws_kpis.column_dimensions["D"].width = 14
    ws_kpis.column_dimensions["E"].width = 16

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
